import io
import json
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app import ingest
from app.main import create_app
from app.settings import ENV_VARS


def message(content=None, tool_calls=None):
    return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=content, tool_calls=tool_calls))])


class FakeAI:
    """Answers describe_table with a fixed description, as a tool call or as JSON text."""

    def __init__(self, answer: dict):
        self.answer = answer
        self.calls = []
        self.chat = SimpleNamespace(completions=SimpleNamespace(create=self.create))

    def create(self, **kwargs):
        self.calls.append(kwargs)
        text = json.dumps(self.answer)
        if "tools" in kwargs:
            call = SimpleNamespace(function=SimpleNamespace(name="describe_table", arguments=text))
            return message(tool_calls=[call])
        return message(content=f"Here you go:\n```json\n{text}\n```")


class FakeFetch:
    def __init__(self, body=b"a,b\n1,2\n", content_type="text/csv", filename="Budget - Sheet1.csv"):
        self.response = (body, content_type, filename)
        self.urls = []

    def __call__(self, url):
        self.urls.append(url)
        return self.response


@pytest.fixture(autouse=True)
def clean_env(monkeypatch):
    for var in ENV_VARS.values():
        monkeypatch.delenv(var, raising=False)


def make(tmp_path, ai=None, fetch=None):
    app = create_app(data_dir=tmp_path, client_factory=lambda s: ai, fetch_url=fetch or FakeFetch())
    return TestClient(app)


def new_project(api, name="Demo"):
    return api.post("/api/projects", json={"name": name}).json()["id"]


def upload(api, pid, *files):
    return api.post(f"/api/projects/{pid}/sources/upload", files=[("files", f) for f in files])


def xlsx(sheets: dict[str, list[list]]) -> bytes:
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def column(profile, name):
    return next(c for c in profile["columns"] if c["name"] == name)


# Reading files


def test_csv_delimiter_blanks_and_long_rows():
    table = ingest.read_csv("t", b"a;b;c\n1;;x\n\n2;3;y;extra\n")
    assert list(table.frame.columns) == ["a", "b", "c"]
    assert table.frame["a"].tolist() == ["1", "2"] and table.frame["c"].tolist() == ["x", "y"]
    assert table.frame["b"].isna().tolist() == [True, False]
    assert "1 row had more values" in table.warnings[0]


def test_csv_falls_back_to_windows_encoding():
    table = ingest.read_csv("t", "name\nCafé\n".encode("cp1252"))
    assert table.frame["name"][0] == "Café"


def test_headers_are_filled_in_and_made_unique():
    table = ingest.read_csv("t", b"Region,,region,Version 2.0\nN,1,S,x\n")
    assert list(table.frame.columns) == ["Region", "column_2", "region_2", "Version 2.0"]


def test_excel_reads_each_data_sheet_and_skips_notes():
    data = xlsx({"Jan": [["id", "amount"], [1, 10]], "Feb": [["id", "amount"], [2, 20]], "Notes": [["read me"]]})
    tables = ingest.read_file("Sales.xlsx", data)
    assert [t.name for t in tables] == ["Sales Jan", "Sales Feb"]
    assert tables[1].frame.to_dict("records") == [{"id": "2", "amount": "20"}]


def test_single_sheet_workbook_is_named_after_the_file():
    tables = ingest.read_file("Customers.xlsx", xlsx({"Sheet1": [["id"], ["C1"]], "Notes": [["x"]]}))
    assert [t.name for t in tables] == ["Customers"]


def test_unsupported_files_are_rejected():
    with pytest.raises(ingest.IngestError, match="xlsx"):
        ingest.read_file("old.xls", b"")
    with pytest.raises(ingest.IngestError, match="only CSV and Excel"):
        ingest.read_file("notes.pdf", b"")


# Google Sheets


@pytest.mark.parametrize(
    "link, expected",
    [
        (
            "https://docs.google.com/spreadsheets/d/abc_123/edit#gid=42",
            "https://docs.google.com/spreadsheets/d/abc_123/export?format=csv&gid=42",
        ),
        (
            "https://docs.google.com/spreadsheets/d/abc_123/edit?usp=sharing",
            "https://docs.google.com/spreadsheets/d/abc_123/export?format=csv",
        ),
        (
            "https://docs.google.com/spreadsheets/d/e/2PACX-xyz/pubhtml",
            "https://docs.google.com/spreadsheets/d/e/2PACX-xyz/pub?output=csv",
        ),
    ],
)
def test_gsheet_links_become_csv_exports(link, expected):
    assert ingest.gsheet_csv_url(link) == expected


@pytest.mark.parametrize(
    "link",
    ["http://docs.google.com/spreadsheets/d/abc", "https://evil.example/spreadsheets/d/abc", "https://docs.google.com/document/d/abc"],
)
def test_gsheet_rejects_other_hosts_and_documents(link):
    with pytest.raises(ingest.IngestError):
        ingest.gsheet_csv_url(link)


def test_gsheet_source_takes_the_sheet_title(tmp_path):
    fetch = FakeFetch()
    api = make(tmp_path, fetch=fetch)
    pid = new_project(api)
    body = api.post(f"/api/projects/{pid}/sources/gsheet", json={"url": "https://docs.google.com/spreadsheets/d/abc/edit"}).json()
    assert body["added"][0]["table"] == "budget_sheet1"
    assert body["added"][0]["origin"]["kind"] == "gsheet"
    assert fetch.urls == ["https://docs.google.com/spreadsheets/d/abc/export?format=csv"]


def test_private_gsheet_gives_a_sharing_hint(tmp_path):
    api = make(tmp_path, fetch=FakeFetch(b"<html>Sign in</html>", "text/html", ""))
    pid = new_project(api)
    response = api.post(f"/api/projects/{pid}/sources/gsheet", json={"url": "https://docs.google.com/spreadsheets/d/abc"})
    assert response.status_code == 400
    assert "Anyone with the link" in response.json()["detail"]


# Profiling


def test_profile_detects_types_and_issues(tmp_path):
    api = make(tmp_path)
    pid = new_project(api)
    csv = (
        "id,qty,price,day,region,flag\n"
        "1,2,RM 1,200.50,2024-01-05,Central,yes\n"
        "2,3,15.00,05/02/2024,central ,no\n"
        "3,x,N/A,2024-03-01,,yes\n"
        "3,x,N/A,2024-03-01,,yes\n"
    )
    # The quoted price keeps its thousands separator inside one field.
    csv = csv.replace("RM 1,200.50", '"RM 1,200.50"')
    sid = upload(api, pid, ("orders.csv", csv.encode())).json()["added"][0]["id"]
    profile = api.get(f"/api/projects/{pid}/sources/{sid}").json()["profile"]

    assert profile["rows"] == 4 and profile["duplicate_rows"] == 1
    assert column(profile, "id")["detected_type"] == "integer"
    assert column(profile, "id")["unique"] is False
    qty = column(profile, "qty")
    assert qty["detected_type"] == "text"  # half the values are "x"
    price = column(profile, "price")
    assert price["detected_type"] == "decimal" and price["max"] == 1200.5 and price["placeholders"] == 2
    assert any("currency symbols" in i for i in price["issues"])
    day = column(profile, "day")
    assert day["detected_type"] == "date" and any("2 different formats" in i for i in day["issues"])
    region = column(profile, "region")
    assert region["blanks"] == 2
    assert any("spaces" in i for i in region["issues"]) and any("case" in i for i in region["issues"])
    assert column(profile, "flag")["detected_type"] == "boolean"


def test_mostly_numeric_column_reports_mixed_types(tmp_path):
    api = make(tmp_path)
    pid = new_project(api)
    values = "\n".join(["n"] + [str(i) for i in range(9)] + ["ten"])
    sid = upload(api, pid, ("t.csv", values.encode())).json()["added"][0]["id"]
    n = column(api.get(f"/api/projects/{pid}/sources/{sid}").json()["profile"], "n")
    assert n["detected_type"] == "integer" and n["type_confidence"] == 0.9
    assert any("mixed types" in i for i in n["issues"])


# Projects and sources API


def test_project_lifecycle(tmp_path):
    api = make(tmp_path)
    pid = new_project(api, "Sales review")
    assert [p["name"] for p in api.get("/api/projects").json()] == ["Sales review"]
    assert api.patch(f"/api/projects/{pid}", json={"name": "Q3 review"}).json()["name"] == "Q3 review"
    assert api.get(f"/api/projects/{pid}").json()["workspace_id"] == "local"
    assert api.delete(f"/api/projects/{pid}").status_code == 200
    assert api.get(f"/api/projects/{pid}").status_code == 404


def test_bad_project_ids_are_not_found(tmp_path):
    api = make(tmp_path)
    assert api.get("/api/projects/..%2F..%2Fsettings").status_code == 404
    assert api.get("/api/projects/000000000000").status_code == 404


def test_upload_preview_rename_and_delete_source(tmp_path):
    api = make(tmp_path)
    pid = new_project(api)
    body = upload(api, pid, ("Sales 2024.csv", b"a,b\n1,2\n3,4\n"), ("Sales 2024.csv", b"a\n9\n")).json()
    first, second = body["added"]
    assert (first["table"], second["table"]) == ("sales_2024", "sales_2024_2")
    assert first["rows"] == 2 and first["columns"] == 2

    preview = api.get(f"/api/projects/{pid}/sources/{first['id']}/preview?limit=1&offset=1").json()
    assert preview == {"columns": ["a", "b"], "rows": [["3", "4"]], "total": 2, "offset": 1}

    clash = api.patch(f"/api/projects/{pid}/sources/{first['id']}", json={"name": "Sales 2024 2"})
    assert clash.status_code == 400
    renamed = api.patch(f"/api/projects/{pid}/sources/{first['id']}", json={"name": "Sales"}).json()
    assert renamed["table"] == "sales"
    assert api.get(f"/api/projects/{pid}/sources/{first['id']}/preview").json()["total"] == 2

    api.delete(f"/api/projects/{pid}/sources/{second['id']}")
    assert [s["table"] for s in api.get(f"/api/projects/{pid}").json()["sources"]] == ["sales"]


def test_upload_reports_bad_files_and_keeps_good_ones(tmp_path):
    api = make(tmp_path)
    pid = new_project(api)
    body = upload(api, pid, ("ok.csv", b"a\n1\n"), ("bad.pdf", b"%PDF")).json()
    assert len(body["added"]) == 1 and "only CSV and Excel" in body["errors"][0]
    assert upload(api, pid, ("bad.pdf", b"%PDF")).status_code == 400


def test_sample_dataset_loads(tmp_path):
    api = make(tmp_path)
    pid = new_project(api)
    tables = {s["table"]: s for s in api.post(f"/api/projects/{pid}/sources/sample").json()["added"]}
    assert set(tables) == {"customers", "products", "sales"}
    assert tables["sales"]["issue_count"] > 0


# AI profiling

ANSWER = {
    "summary": "Orders by customer.",
    "grain": "one order",
    "table_role": "fact",
    "columns": [
        {"name": "order_id", "description": "Order number", "role": "id", "suggested_type": "text", "issues": ["None."]},
        {"name": "amount", "description": "Order value", "role": "measure", "suggested_type": "decimal", "issues": ["RM prefix"]},
        {"name": "invented", "description": "Not a real column", "role": "id", "suggested_type": "text"},
    ],
}


def setup_source(tmp_path, ai, capabilities):
    api = make(tmp_path, ai=ai)
    api.put("/api/settings", json={"api_key": "sk-secret-value-1234"})
    settings = json.loads((tmp_path / "settings.json").read_text())
    settings["capabilities"] = capabilities
    (tmp_path / "settings.json").write_text(json.dumps(settings))
    pid = new_project(api)
    sid = upload(api, pid, ("orders.csv", b"order_id,amount\nA1,RM 10\nA2,20\n")).json()["added"][0]["id"]
    return api, pid, sid


@pytest.mark.parametrize("capabilities", [{"tools": True, "json_mode": True}, {"tools": False, "json_mode": True}, {}])
def test_describe_with_each_protocol(tmp_path, capabilities):
    ai = FakeAI(ANSWER)
    api, pid, sid = setup_source(tmp_path, ai, capabilities)
    source = api.post(f"/api/projects/{pid}/sources/{sid}/describe").json()
    assert source["ai"]["table_role"] == "fact" and source["ai"]["status"] == "AI"
    assert set(source["ai"]["columns"]) == {"order_id", "amount"}  # the invented column is dropped
    assert source["ai"]["columns"]["amount"]["issues"] == ["RM prefix"]
    assert source["ai"]["columns"]["order_id"]["issues"] == []  # "None." is not an issue
    call = ai.calls[0]
    assert ("tools" in call) == bool(capabilities.get("tools"))
    assert ("response_format" in call) == bool(capabilities.get("json_mode") and not capabilities.get("tools"))


def test_describe_sends_only_profiles_and_a_few_rows(tmp_path):
    ai = FakeAI(ANSWER)
    api = make(tmp_path, ai=ai)
    api.put("/api/settings", json={"api_key": "sk-secret-value-1234"})
    pid = new_project(api)
    rows = "\n".join(["id"] + [f"row-{i}" for i in range(50)])
    sid = upload(api, pid, ("t.csv", rows.encode())).json()["added"][0]["id"]
    api.post(f"/api/projects/{pid}/sources/{sid}/describe")
    prompt = ai.calls[0]["messages"][-1]["content"]
    assert "row-4" in prompt and "row-30" not in prompt


def test_describe_needs_a_key(tmp_path):
    api = make(tmp_path, ai=FakeAI(ANSWER))
    pid = new_project(api)
    sid = upload(api, pid, ("t.csv", b"a\n1\n")).json()["added"][0]["id"]
    assert api.post(f"/api/projects/{pid}/sources/{sid}/describe").status_code == 400


def test_reprofile_clears_the_ai_description(tmp_path):
    api, pid, sid = setup_source(tmp_path, FakeAI(ANSWER), {"tools": True})
    api.post(f"/api/projects/{pid}/sources/{sid}/describe")
    assert api.post(f"/api/projects/{pid}/sources/{sid}/profile").json()["ai"] is None


def test_sql_cannot_reach_outside_files(tmp_path):
    from app.projects import ProjectStore

    store = ProjectStore(tmp_path)
    project = store.create("Locked")
    with store.connect(project["id"]) as con:
        with pytest.raises(Exception, match="(?i)external|permission"):
            con.execute(f"SELECT * FROM read_csv('{(tmp_path / 'x.csv').as_posix()}')")
